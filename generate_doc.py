import docx
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.oxml import OxmlElement, parse_xml
from docx.oxml.ns import nsdecls, qn
import os

def set_cell_background(cell, fill_hex):
    tcPr = cell._element.get_or_add_tcPr()
    shd = parse_xml(f'<w:shd {nsdecls("w")} w:fill="{fill_hex}"/>')
    tcPr.append(shd)

def set_cell_margins(cell, top=100, bottom=100, left=150, right=150):
    tcPr = cell._element.get_or_add_tcPr()
    tcMar = parse_xml(f'<w:tcMar {nsdecls("w")}><w:top w:w="{top}" w:type="dxa"/><w:bottom w:w="{bottom}" w:type="dxa"/><w:left w:w="{left}" w:type="dxa"/><w:right w:w="{right}" w:type="dxa"/></w:tcMar>')
    tcPr.append(tcMar)

def add_header_footer(doc):
    for section in doc.sections:
        # Header
        header = section.header
        hp = header.paragraphs[0]
        hp.alignment = WD_ALIGN_PARAGRAPH.RIGHT
        hrun = hp.add_run("ST6003CEM Web API Development")
        hrun.font.name = "Arial"
        hrun.font.size = Pt(9)
        hrun.font.color.rgb = RGBColor(120, 120, 120)

        # Footer
        footer = section.footer
        fp = footer.paragraphs[0]
        fp.alignment = WD_ALIGN_PARAGRAPH.LEFT
        frun = fp.add_run("Archive Outfitters | Sabja Shrestha (15373146)")
        frun.font.name = "Arial"
        frun.font.size = Pt(9)
        frun.font.color.rgb = RGBColor(120, 120, 120)

def build_document():
    doc = docx.Document()

    # Set Margins (1 inch)
    for section in doc.sections:
        section.top_margin = Inches(1)
        section.bottom_margin = Inches(1)
        section.left_margin = Inches(1)
        section.right_margin = Inches(1)

    # -------------------------------------------------------------
    # COVER PAGE
    # -------------------------------------------------------------
    p_institution = doc.add_paragraph()
    p_institution.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r_inst = p_institution.add_run("Softwarica College of IT & E-Commerce\nin collaboration with Coventry University\n\n")
    r_inst.font.name = "Arial"
    r_inst.font.size = Pt(14)
    r_inst.font.bold = True
    r_inst.font.color.rgb = RGBColor(20, 40, 80)

    p_title = doc.add_paragraph()
    p_title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r_sub = p_title.add_run("A Project Report On\n")
    r_sub.font.name = "Arial"
    r_sub.font.size = Pt(16)
    r_sub.font.bold = True
    
    r_main = p_title.add_run("Archive Outfitters: A Community-Driven Vintage & Streetwear E-Commerce Platform and RESTful Web API\n\n")
    r_main.font.name = "Arial"
    r_main.font.size = Pt(20)
    r_main.font.bold = True
    r_main.font.color.rgb = RGBColor(10, 80, 160)

    p_meta = doc.add_paragraph()
    p_meta.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r_meta = p_meta.add_run(
        "ST6003CEM Web API Development\n"
        "BSc (Hons) in Computing\n\n\n"
        "Submitted By: Sabja Shrestha\n"
        "Coventry ID: 15373146\n"
        "Batch: 35C\n"
        "Module Leader: Mr. Albert Maharjan\n"
    )
    r_meta.font.name = "Arial"
    r_meta.font.size = Pt(12)
    r_meta.font.bold = True
    r_meta.font.color.rgb = RGBColor(50, 50, 50)

    doc.add_page_break()

    # Add header/footer after cover page section setup
    add_header_footer(doc)

    # Read the markdown content
    md_path = "ST6003CEM_Web_API_Development_Report.md"
    if not os.path.exists(md_path):
        print("Markdown file not found!")
        return

    with open(md_path, "r", encoding="utf-8") as f:
        lines = f.readlines()

    in_code_block = False
    code_lines = []

    for line in lines:
        line_str = line.strip()

        # Skip main cover title lines from body render
        if line_str.startswith("# A Project Report On") or line_str.startswith("## Archive Outfitters"):
            continue
        if "Submitted By: Sabja Shrestha" in line_str or "Coventry ID: 15373146" in line_str:
            continue

        # Code block handler
        if line_str.startswith("```"):
            if in_code_block:
                # End code block
                in_code_block = False
                code_text = "".join(code_lines)
                tbl = doc.add_table(rows=1, cols=1)
                tbl.alignment = WD_TABLE_ALIGNMENT.CENTER
                cell = tbl.cell(0, 0)
                set_cell_background(cell, "F4F6F8")
                set_cell_margins(cell, top=120, bottom=120, left=180, right=180)
                cp = cell.paragraphs[0]
                cp.paragraph_format.space_before = Pt(4)
                cp.paragraph_format.space_after = Pt(4)
                crun = cp.add_run(code_text)
                crun.font.name = "Consolas"
                crun.font.size = Pt(9.5)
                crun.font.color.rgb = RGBColor(30, 30, 30)
                code_lines = []
                p_spacer = doc.add_paragraph()
                p_spacer.paragraph_format.space_after = Pt(6)
            else:
                in_code_block = True
                code_lines = []
            continue

        if in_code_block:
            code_lines.append(line)
            continue

        # Headings
        if line_str.startswith("## "):
            h_text = line_str[3:]
            h = doc.add_heading(h_text, level=1)
            h.paragraph_format.space_before = Pt(16)
            h.paragraph_format.space_after = Pt(6)
            for r in h.runs:
                r.font.name = "Arial"
                r.font.size = Pt(16)
                r.font.bold = True
                r.font.color.rgb = RGBColor(10, 70, 140)
        elif line_str.startswith("### "):
            h_text = line_str[4:]
            h = doc.add_heading(h_text, level=2)
            h.paragraph_format.space_before = Pt(12)
            h.paragraph_format.space_after = Pt(4)
            for r in h.runs:
                r.font.name = "Arial"
                r.font.size = Pt(13)
                r.font.bold = True
                r.font.color.rgb = RGBColor(40, 90, 160)
        elif line_str.startswith("#### "):
            h_text = line_str[5:]
            h = doc.add_heading(h_text, level=3)
            h.paragraph_format.space_before = Pt(8)
            h.paragraph_format.space_after = Pt(2)
            for r in h.runs:
                r.font.name = "Arial"
                r.font.size = Pt(11)
                r.font.bold = True
        elif line_str.startswith("- ") or line_str.startswith("* "):
            p = doc.add_paragraph(style='List Bullet')
            p.paragraph_format.space_after = Pt(3)
            r = p.add_run(line_str[2:])
            r.font.name = "Arial"
            r.font.size = Pt(10.5)
        elif line_str.startswith("*Vis ") or line_str.startswith("_Vis "):
            # Figure Caption
            p = doc.add_paragraph()
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            p.paragraph_format.space_before = Pt(4)
            p.paragraph_format.space_after = Pt(10)
            r = p.add_run(line_str.strip("*_"))
            r.font.name = "Arial"
            r.font.size = Pt(9.5)
            r.font.italic = True
            r.font.color.rgb = RGBColor(90, 90, 90)
        elif line_str == "---":
            doc.add_page_break()
        elif line_str:
            p = doc.add_paragraph()
            p.paragraph_format.space_after = Pt(6)
            p.paragraph_format.line_spacing = 1.15
            r = p.add_run(line_str)
            r.font.name = "Arial"
            r.font.size = Pt(10.5)
            r.font.color.rgb = RGBColor(30, 30, 30)

    output_docx = "ST6003CEM_Web_API_Development_Report.docx"
    doc.save(output_docx)
    print(f"Successfully generated {output_docx}")

if __name__ == "__main__":
    build_document()
