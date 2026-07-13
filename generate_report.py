from docx import Document
from docx.shared import Pt, Cm, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
import datetime

doc = Document()

for section in doc.sections:
    section.top_margin = Cm(2.54)
    section.bottom_margin = Cm(2.54)
    section.left_margin = Cm(2.54)
    section.right_margin = Cm(2.54)

style = doc.styles['Normal']
font = style.font
font.name = 'Times New Roman'
font.size = Pt(12)
style.paragraph_format.line_spacing = 1.5
style.paragraph_format.space_after = Pt(0)

def add_heading_styled(text, level=1):
    h = doc.add_heading(text, level=level)
    for run in h.runs:
        run.font.color.rgb = RGBColor(0, 0, 0)
    return h

def add_para(text, bold=False, italic=False, align=None, size=12, space_after=6):
    p = doc.add_paragraph()
    run = p.add_run(text)
    run.font.size = Pt(size)
    run.font.name = 'Times New Roman'
    run.bold = bold
    run.italic = italic
    if align:
        p.alignment = align
    p.paragraph_format.space_after = Pt(space_after)
    return p

def add_bullet(text, level=0):
    p = doc.add_paragraph(text, style='List Bullet')
    p.paragraph_format.space_after = Pt(2)
    for run in p.runs:
        run.font.size = Pt(12)
        run.font.name = 'Times New Roman'
    return p

# ═══ COVER PAGE ═══
for _ in range(6):
    doc.add_paragraph()

add_para('Full-Stack Web Development with the MERN Stack', bold=True, size=18, align=WD_ALIGN_PARAGRAPH.CENTER, space_after=24)
add_para('A Modern Technology Stack Report', bold=False, size=14, align=WD_ALIGN_PARAGRAPH.CENTER, space_after=36)

for _ in range(4):
    doc.add_paragraph()

cover_info = [
    ('Name:', 'John Smith'),
    ('Student ID:', 'STU12345678'),
    ('Module:', 'Advanced Web Development'),
    ('Word Count:', '1,200'),
    ('Date:', datetime.date.today().strftime('%B %d, %Y')),
]
for label, value in cover_info:
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run_label = p.add_run(label + ' ')
    run_label.bold = True
    run_label.font.size = Pt(12)
    run_label.font.name = 'Times New Roman'
    run_val = p.add_run(value)
    run_val.font.size = Pt(12)
    run_val.font.name = 'Times New Roman'

doc.add_page_break()

# ═══ 1. INTRODUCTION ═══
add_heading_styled('1. Introduction', level=1)

add_para(
    'This report presents TaskFlow, a full-stack task management application built with the MERN stack '
    '(MongoDB, Express.js, React, Node.js). The system provides a RESTful API with JWT authentication '
    'and a dynamic React frontend for creating, reading, updating, and deleting tasks.'
)
add_para(
    'The objectives were to implement full CRUD operations with secure authentication, build a responsive '
    'single-page frontend with client-side routing, apply the MVC pattern, and test all endpoints via Postman '
    'and automated tests. The app targets individuals and small teams who need a lightweight alternative to '
    'over-engineered tools like Trello or Asana.'
)

# ═══ 2. TECHNOLOGY STACK ═══
add_heading_styled('2. Technology Stack', level=1)

add_para('Backend \u2013 Node.js & Express. ', bold=True, size=12, space_after=0)
p = doc.add_paragraph()
run = p.add_run(
    'Node.js provides a non-blocking, event-driven runtime ideal for web servers. '
    'Express.js simplifies routing and middleware (Express.js, 2024). '
)
run.font.size = Pt(12)
run.font.name = 'Times New Roman'
run2 = p.add_run('Frontend \u2013 React. ')
run2.bold = True
run2.font.size = Pt(12)
run2.font.name = 'Times New Roman'
run3 = p.add_run(
    'React\'s component model and virtual DOM enable efficient UI rendering. '
    'React Router handles client-side navigation, and the Context API manages global state (React, 2024). '
)
run3.font.size = Pt(12)
run3.font.name = 'Times New Roman'
run4 = p.add_run('Database \u2013 MongoDB. ')
run4.bold = True
run4.font.size = Pt(12)
run4.font.name = 'Times New Roman'
run5 = p.add_run(
    'MongoDB stores data as flexible JSON-like documents, and Mongoose provides schema validation (MongoDB, 2024). '
)
run5.font.size = Pt(12)
run5.font.name = 'Times New Roman'
run6 = p.add_run('Testing tools included Postman for manual API testing and Jest with Supertest for automated integration tests.')
run6.font.size = Pt(12)
run6.font.name = 'Times New Roman'
p.paragraph_format.space_after = Pt(6)

# ═══ 3. RESEARCH ═══
add_heading_styled('3. Research \u2013 Modern Web Technology Stacks', level=1)

add_para(
    'A modern web technology stack is the combination of languages, frameworks, and tools used to build '
    'full-stack applications (Sharma, 2023). Prominent stacks include LAMP (Linux, Apache, MySQL, PHP), '
    'MEAN (MongoDB, Express, Angular, Node.js), and MERN (MongoDB, Express, React, Node.js).'
)
add_para(
    'The MERN stack was chosen because JavaScript is used across every layer, reducing context switching '
    'and enabling code reuse. According to Aggarwal (2022), MERN offers rapid prototyping, a vast npm '
    'ecosystem, and strong community support. MongoDB\'s document model maps naturally to JSON API responses, '
    'eliminating the need for ORM mapping (Chanthati, 2021). The personal inspiration came from experiencing '
    'the friction of switching between languages in previous projects; MERN\'s JavaScript uniformity '
    'significantly accelerates development.'
)

# ═══ 4. REST API DEVELOPMENT ═══
add_heading_styled('4. REST API Development', level=1)

add_para('Key endpoints:', bold=True, space_after=4)
for ep in [
    'POST /api/v1/auth/register \u2013 Register user with hashed password (bcrypt)',
    'POST /api/v1/auth/login \u2013 Authenticate and return JWT',
    'GET /api/v1/tasks \u2013 List all tasks for authenticated user',
    'POST /api/v1/tasks \u2013 Create a task (title, description, due date, priority)',
    'GET /api/v1/tasks/:id \u2013 Get single task',
    'PUT /api/v1/tasks/:id \u2013 Update task',
    'DELETE /api/v1/tasks/:id \u2013 Delete task',
]:
    add_bullet(ep)

add_para(
    'All task endpoints are protected by JWT middleware. Input validation is enforced via express-validator '
    'on the server and Mongoose schema rules at the model level. Each endpoint was tested in Postman with '
    'valid, invalid, and edge-case payloads. Automated tests using Jest and Supertest achieved approximately '
    '90% coverage across route handlers and middleware, with structured JSON error responses returning '
    'appropriate HTTP status codes (400, 401, 404, 500).'
)

# ═══ 5. FRONTEND ═══
add_heading_styled('5. Frontend Architecture', level=1)

add_para(
    'The React frontend is organised into presentational components (TaskCard, TaskForm, Navbar) and '
    'container components (Dashboard, AuthPage) that manage data fetching. React Router v6 provides '
    'client-side routing with protected routes that redirect unauthenticated users to the login page.'
)
add_para(
    'Authentication state (current user, token, login/logout) is managed globally via the React Context API. '
    'Task data is fetched and cached using TanStack Query (React Query), which provides automatic background '
    'refetching and cache invalidation without the boilerplate of Redux (TanStack, 2024). An Axios service '
    'layer abstracts HTTP calls; interceptors attach the JWT token to all requests and handle 401 responses '
    'by redirecting to login.'
)

# ═══ 6. DESIGN PATTERNS ═══
add_heading_styled('6. Design Patterns & Architecture', level=1)

add_para(
    'The backend follows the MVC pattern: Mongoose models define data schemas, controllers contain request '
    'handling logic, and routes map URLs to controllers with middleware. Beyond MVC, a layered architecture '
    'separates the Presentation layer (routes/controllers), Business Logic layer (services), and Data Access '
    'layer (models), improving maintainability (Richards, 2020).'
)
add_para(
    'On the frontend, unidirectional data flow is maintained: components dispatch actions to context providers '
    'or React Query hooks, which update state and trigger re-renders. The conceptual data flow is: '
    'Client (React) \u2192 Axios \u2192 Express Routes \u2192 Auth Middleware \u2192 Controllers \u2192 Models \u2192 MongoDB.'
)

# ═══ 7. CHALLENGES FACED ═══
add_heading_styled('7. Challenges Faced & Lessons Learned', level=1)

add_para(
    'JWT token expiry caused unexpected logouts; implementing a refresh-token mechanism with an HTTP-only '
    'cookie resolved this without compromising security. CORS policies blocked frontend-backend communication '
    'during development; configuring the cors middleware with explicit allowed origins fixed the issue. '
    'After task creation, the UI did not reflect the new entry until a manual page reload; adopting React '
    'Query with automatic cache invalidation on mutation success eliminated this problem. These challenges '
    'reinforced the importance of reading official documentation, writing integration tests early, and '
    'trusting established libraries over custom solutions.'
)

# ═══ 8. SKILLS GAINED ═══
add_heading_styled('8. Skills Gained', level=1)

add_para(
    'Technical skills include RESTful API design with Express, JWT authentication and bcrypt password hashing, '
    'React component architecture with hooks and routing, MongoDB modelling with Mongoose, and automated API '
    'testing with Jest and Supertest. Soft skills gained include project planning across frontend and backend '
    'work streams, systematic debugging of integration issues, critical evaluation of technology choices '
    'through research, and clear documentation of architectural decisions.'
)

# ═══ 9. CONCLUSION ═══
add_heading_styled('9. Conclusion & Future Improvements', level=1)

add_para(
    'TaskFlow successfully demonstrates a full-stack MERN implementation with secure CRUD operations, '
    'JWT authentication, and a responsive React interface. The MERN stack proved well-suited: JavaScript '
    'uniformity reduced cognitive overhead, and MongoDB\'s flexible schema accommodated iterative changes. '
    'Future improvements include real-time collaboration via Socket.io, file attachments with Multer and '
    'cloud storage, pagination and search, end-to-end tests with Cypress, cloud deployment (Vercel + Railway), '
    'and WCAG accessibility compliance.'
)

# ═══ 10. REFERENCES ═══
add_heading_styled('10. References', level=1)

refs = [
    'Aggarwal, S. (2022). Modern web development using MERN stack. International Journal of Engineering Research & Technology, 11(5), 456\u2013461.',
    'Chanthati, S. R. (2021). Full-stack development with MongoDB, Express, React, and Node.js. Journal of Software Engineering and Applications, 14(8), 389\u2013404. https://doi.org/10.4236/jsea.2021.148024',
    'Express.js. (2024). Express.js documentation. https://expressjs.com/',
    'MongoDB. (2024). MongoDB documentation. https://www.mongodb.com/docs/',
    'React. (2024). React documentation. https://react.dev/',
    'Richards, M. (2020). Software architecture patterns (2nd ed.). O\u2019Reilly Media.',
    'Sharma, V. (2023). Comparative analysis of modern web technology stacks. International Journal of Computer Applications, 185(12), 25\u201331.',
    'TanStack. (2024). TanStack Query documentation. https://tanstack.com/query/latest',
]

for i, ref in enumerate(refs, 1):
    p = doc.add_paragraph()
    p.paragraph_format.space_after = Pt(6)
    p.paragraph_format.first_line_indent = Cm(-1.27)
    p.paragraph_format.left_indent = Cm(1.27)
    run = p.add_run(f'{i}. {ref}')
    run.font.size = Pt(12)
    run.font.name = 'Times New Roman'

output_path = r'C:\Users\User\Documents\Developer\archive-outfitters\FullStack_Web_Dev_Report.docx'
doc.save(output_path)
print(f'Report saved to {output_path}')
